"""Reading an uploaded workbook (the registry import, 2026-10-14).

The file comes from outside, so what it may be is narrow: an .xlsx whose first
row is the headings. Headings are matched loosely - any order, any case, the
template's '*' and hints - and anything else is refused with a sentence a
person can act on.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from openpyxl import Workbook

from cmp.core.errors import ValidationFailed
from cmp.infrastructure import spreadsheets
from cmp.infrastructure.spreadsheets import Column, SheetSpec, build_workbook, read_rows

COLUMNS = (
    Column("code", "Source code", required=True),
    Column("name", "Name", required=True),
    Column("role", "Role", choices=("identity", "collection")),
)


def _book(*rows: tuple[Any, ...], title: str = "Sheet") -> bytes:
    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = title
    for row in rows:
        ws.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_headings_match_in_any_order_and_case_with_the_templates_marks() -> None:
    payload = _book(("NAME *", "role", "Source Code (unique)"), ("Studio", "collection", " S-1 "))
    assert read_rows(payload, columns=COLUMNS) == [
        (2, {"name": "Studio", "role": "collection", "code": "S-1"})
    ]


def test_blank_rows_are_skipped_and_row_numbers_are_the_spreadsheets() -> None:
    payload = _book(("Source code", "Name"), ("A", "One"), (None, None), ("", "  "), ("B", "Two"))
    assert [n for n, _ in read_rows(payload, columns=COLUMNS)] == [2, 5]


def test_a_missing_required_column_is_named() -> None:
    with pytest.raises(ValidationFailed, match="missing the column\\(s\\): Name"):
        read_rows(_book(("Source code", "Role"), ("A", "identity")), columns=COLUMNS)


def test_the_named_sheet_is_read_before_the_first() -> None:
    book = Workbook()
    first = book.active
    assert first is not None
    first.append(["Notes"])
    second = book.create_sheet("Data sources")
    second.append(["Source code", "Name"])
    second.append(["A", "One"])
    out = io.BytesIO()
    book.save(out)
    rows = read_rows(out.getvalue(), columns=COLUMNS, sheet="data sources")
    assert rows == [(2, {"code": "A", "name": "One"})]


@pytest.mark.parametrize(
    "payload",
    [b"source_code,name\nA,One\n", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1old xls", b"PK\x03\x04junk"],
)
def test_anything_but_an_xlsx_is_refused_plainly(payload: bytes) -> None:
    with pytest.raises(ValidationFailed) as caught:
        read_rows(payload, columns=COLUMNS)
    assert caught.value.field == "file"


def test_more_than_the_limit_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(spreadsheets, "MAX_ROWS", 2)
    payload = _book(("Source code", "Name"), ("A", "1"), ("B", "2"), ("C", "3"))
    with pytest.raises(ValidationFailed, match="more than 2 rows"):
        read_rows(payload, columns=COLUMNS)


def test_a_formula_is_never_read_as_one() -> None:
    payload = _book(("Source code", "Name"), ("A", '=HYPERLINK("http://x","y")'))
    # Never computed by Excel, so it has no cached value: it reads as empty,
    # and a required cell left empty is the importer's to report.
    assert read_rows(payload, columns=COLUMNS) == [(2, {"code": "A", "name": None})]


def test_a_template_reads_back_as_its_example_rows() -> None:
    spec = SheetSpec("Data sources", COLUMNS, examples=[("A", "One", "identity")])
    payload = build_workbook([spec], instructions=["How to fill"])
    assert read_rows(payload, columns=COLUMNS, sheet="Data sources") == [
        (2, {"code": "A", "name": "One", "role": "identity"})
    ]
