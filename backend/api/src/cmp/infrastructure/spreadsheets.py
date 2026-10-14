"""Excel workbooks in and out (2026-10-14): the registry imports.

Reading is defensive, because the file comes from outside:

* **Only .xlsx.** It is a zip of XML; anything that does not open as one is
  refused with a plain message. The old binary `.xls` is not read.
* **No XML tricks.** `defusedxml` is installed, and `openpyxl` uses it when it
  is, so an entity-expansion bomb is refused rather than parsed.
* **Values, not formulas.** The workbook is read with `data_only`, so a cell
  holding `=HYPERLINK(...)` gives its last computed value, never the formula.
* **Bounded.** At most `MAX_ROWS` data rows, and only the columns a caller
  names; a cell's text is stripped and capped.

Writing builds a template: a header row, optional example rows, a dropdown on
each column with a fixed set of values, and a sheet of instructions.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from cmp.core.errors import ValidationFailed

#: More than this is a migration, not an import.
MAX_ROWS = 1000
#: No registry value needs more; a longer cell is a mistake or an attack.
MAX_CELL_CHARS = 2000

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True, slots=True)
class Column:
    """One column of a template, and how its header is recognised."""

    key: str
    header: str
    required: bool = False
    #: Shown in the instructions sheet.
    help: str = ""
    #: A fixed set of values: a dropdown in the template.
    choices: tuple[str, ...] = ()
    width: int = 24
    #: Other header texts that mean this column.
    aliases: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class SheetSpec:
    title: str
    columns: tuple[Column, ...]
    examples: Sequence[Sequence[Any]] = field(default_factory=tuple)


def _normalise(text: object) -> str:
    """'Legal name *' -> 'legal name'; 'Security confirmed on (YYYY-MM-DD)' -> its words."""
    out = str(text or "").replace("*", " ").lower()
    if "(" in out:
        out = out.split("(", 1)[0]
    return " ".join(out.replace("_", " ").split())


def _cell(value: object) -> object:
    if isinstance(value, str):
        value = value.strip()[:MAX_CELL_CHARS]
        return value or None
    return value


def read_rows(
    payload: bytes, *, columns: Sequence[Column], sheet: str | None = None
) -> list[tuple[int, dict[str, Any]]]:
    """The data rows of one sheet as (spreadsheet row number, {column key: value}).

    The first sheet whose title matches `sheet` (any case), else the first
    sheet. The first row is the header. Rows with nothing in them are skipped.
    """
    if not zipfile.is_zipfile(io.BytesIO(payload)):
        raise ValidationFailed(
            "That is not an Excel workbook (.xlsx). Save the file as .xlsx and upload it again.",
            field="file",
        )
    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # a corrupt or hostile file - say so, plainly
        raise ValidationFailed(
            "The workbook could not be read. Save it again as .xlsx from Excel and upload it.",
            field="file",
        ) from exc
    try:
        names = {name.lower(): name for name in book.sheetnames}
        title = names.get(sheet.lower()) if sheet else None
        ws = book[title] if title else book.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValidationFailed(
                "The sheet is empty: the first row must be the headings.", field="file"
            )

        wanted: dict[str, Column] = {}
        for col in columns:
            for text in (col.header, col.key, *col.aliases):
                wanted[_normalise(text)] = col
        position: dict[int, Column] = {}
        for index, text in enumerate(header):
            match = wanted.get(_normalise(text))
            if match is not None and match not in position.values():
                position[index] = match
        missing = [c.header for c in columns if c.required and c not in position.values()]
        if missing:
            raise ValidationFailed(
                f"The sheet '{ws.title}' is missing the column(s): {', '.join(missing)}. "
                "Download the template to see the headings it expects.",
                field="file",
            )

        out: list[tuple[int, dict[str, Any]]] = []
        for number, values in enumerate(rows, start=2):
            record = {
                col.key: _cell(values[index]) if index < len(values) else None
                for index, col in position.items()
            }
            if all(v is None for v in record.values()):
                continue
            if len(out) >= MAX_ROWS:
                raise ValidationFailed(
                    f"The sheet has more than {MAX_ROWS} rows. Split it into smaller files.",
                    field="file",
                )
            out.append((number, record))
        return out
    finally:
        book.close()


_HEADER_FILL = PatternFill("solid", fgColor="1F5C9E")
_REQUIRED_FILL = PatternFill("solid", fgColor="0E3A6B")


def build_workbook(sheets: Iterable[SheetSpec], *, instructions: Sequence[str] = ()) -> bytes:
    """A template: one sheet per spec, with dropdowns, and an instructions sheet."""
    book = Workbook()
    book.remove(book.active)
    specs = list(sheets)
    for spec in specs:
        ws = book.create_sheet(spec.title)
        for index, col in enumerate(spec.columns, start=1):
            cell = ws.cell(
                row=1, column=index, value=f"{col.header} *" if col.required else col.header
            )
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _REQUIRED_FILL if col.required else _HEADER_FILL
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(index)].width = col.width
            if col.choices:
                letter = get_column_letter(index)
                rule = DataValidation(
                    type="list",
                    formula1='"' + ",".join(col.choices) + '"',
                    allow_blank=not col.required,
                    showErrorMessage=True,
                    errorTitle=col.header,
                    error=f"Choose one of: {', '.join(col.choices)}",
                )
                rule.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
                ws.add_data_validation(rule)
        for r, row in enumerate(spec.examples, start=2):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, date | datetime):
                    cell.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    guide = book.create_sheet("How to fill")
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 90
    line = 1
    for text in instructions:
        guide.cell(row=line, column=1, value=text).font = Font(bold=line == 1)
        line += 1
    for spec in specs:
        line += 1
        guide.cell(row=line, column=1, value=f"Sheet: {spec.title}").font = Font(bold=True, size=12)
        line += 1
        for title, value in (("Column", 1), ("Required", 2), ("What to put in it", 3)):
            guide.cell(row=line, column=value, value=title).font = Font(bold=True)
        line += 1
        for col in spec.columns:
            allowed = f" Allowed: {', '.join(col.choices)}." if col.choices else ""
            guide.cell(row=line, column=1, value=col.header)
            guide.cell(row=line, column=2, value="Yes" if col.required else "No")
            guide.cell(row=line, column=3, value=f"{col.help}{allowed}").alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            line += 1
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
